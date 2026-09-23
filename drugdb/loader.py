"""Идемпотентный загрузчик ГРЛС и открытых данных МДЛП.

Источники и прямые URL намеренно задаются через окружение: портал ГРЛС меняет
URL выгрузок, а Датамаркет публикует версии CSV. Лоадер не зависит от HTML-разметки
портала и может принимать конкретный URL последней опубликованной выгрузки.
"""
import argparse
import csv
import io
import logging
import os
import re
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import psycopg
import requests
from openpyxl import load_workbook

logging.basicConfig(
    level=os.getenv("LOG_LEVEL", "INFO"),
    format="%(asctime)s %(levelname)s [drugdb] %(message)s",
)
LOG = logging.getLogger("drugdb")

DSN = os.getenv("DRUG_DB_DSN", "")
GRLS_URL = os.getenv("GRLS_EXPORT_URL", "").strip()
MDLP_URL = os.getenv(
    "MDLP_EXPORT_URL",
    "https://xn--80aaani3am7aog.xn--80ajghhoc2aj1c8b.xn--p1ai/bi/api/opendata/7731376812-MDLPGtins/data/latest",
).strip()
TIMEOUT = int(os.getenv("DRUGDB_HTTP_TIMEOUT", "180"))
BATCH_SIZE = int(os.getenv("DRUGDB_BATCH_SIZE", "2000"))

HEADER_ALIASES = {
    "reg_number": [
        "номер регистрационного удостоверения", "номер ру", "регистрационный номер",
        "номер регистрационной записи", "reg_number",
    ],
    "trade_name": ["торговое наименование", "наименование лекарственного препарата", "trade_name"],
    "inn": ["мпн", "мнн", "международное непатентованное наименование", "inn"],
    "dosage_form": ["лекарственная форма", "форма выпуска", "dosage_form"],
    "dosage_value": ["дозировка", "дозы", "дозировка/концентрация", "dosage_value"],
    "manufacturer": ["производитель", "наименование производителя", "manufacturer"],
    "holder": [
        "юридическое лицо, на имя которого выдано регистрационное удостоверение",
        "держатель регистрационного удостоверения",
        "держатель/владелец регистрационного удостоверения",
        "владелец регистрационного удостоверения", "holder",
    ],
    "registration_date": ["дата регистрации", "registration_date"],
    "expiry_date": ["дата окончания действия регистрационного удостоверения", "дата окончания действия", "expiry_date"],
    "cancellation_date": ["дата аннулирования регистрационного удостоверения", "дата аннулирования", "cancellation_date"],
    "production_stages": ["сведения о стадиях производства", "стадии производства", "production_stages"],
    "pharmacotherapeutic_group": ["фармако-терапевтическая группа", "фармакотерапевтическая группа", "pharmacotherapeutic_group"],
    "essential_drug": ["наличие лекарственного препарата в перечне жнвлп", "жнвлп", "essential_drug"],
    "contains_controlled_substances": ["наличие в лекарственном препарате наркотических средств, психотропных веществ", "наркотических средств", "controlled_substances"],
    "orphan_status": ["статус признания лекарственного препарата орфанным", "орфанный", "orphan_status"],
    "status": ["состояние", "статус", "status"],
}

GTIN_ALIASES = ["gtin", "гтин", "код gtin", "gtin товара", "гтин товара"]
REG_ALIASES = [
    "номер ру", "номер р.у.", "номер р/у", "№ ру", "номер регистрационного удостоверения",
    "регистрационный номер", "reg_number", "номер регистрации",
    "регистрационное удостоверение", "registration number",
    "registration_number", "ru number", "ru_number",
]
PACKAGE_ALIASES = ["описание упаковки", "описание потребительской упаковки", "упаковка", "характеристика упаковки", "package_desc", "package description"]


def _norm_header(v):
    return re.sub(r"\s+", " ", str(v or "").strip().lower().replace("ё", "е"))


def _find_column(headers, aliases):
    normalized = {_norm_header(h): i for i, h in enumerate(headers)}
    for alias in aliases:
        if _norm_header(alias) in normalized:
            return normalized[_norm_header(alias)]
    for i, h in enumerate(headers):
        nh = _norm_header(h)
        if any(_norm_header(a) in nh for a in aliases):
            return i
    return None


def _text(v, limit=None):
    s = "" if v is None else str(v).strip()
    if limit and len(s) > limit:
        s = s[:limit]
    return s or None


def _normalise_gtin(v):
    digits = re.sub(r"\D", "", str(v or ""))
    if len(digits) == 13:
        digits = "0" + digits
    if len(digits) != 14:
        raise ValueError("GTIN must contain 13 or 14 digits")
    total = 0
    for pos, digit in enumerate(reversed(digits[:-1]), start=1):
        total += int(digit) * (3 if pos % 2 else 1)
    check = (10 - (total % 10)) % 10
    if check != int(digits[-1]):
        raise ValueError("Invalid GTIN check digit")
    return digits


def _download(url, source=None):
    if not url:
        raise ValueError("URL выгрузки не задана")

    headers = {}
    if source and DSN:
        with psycopg.connect(DSN) as conn:
            conn.execute(
                "CREATE TABLE IF NOT EXISTS drug_source_state ("
                "source VARCHAR(32) PRIMARY KEY, url TEXT NOT NULL, etag TEXT, "
                "last_modified TEXT, updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW())"
            )
            row = conn.execute(
                "SELECT url, etag, last_modified FROM drug_source_state WHERE source=%s",
                (source,),
            ).fetchone()
            if row and row[0] == url:
                if row[1]:
                    headers["If-None-Match"] = row[1]
                if row[2]:
                    headers["If-Modified-Since"] = row[2]

    r = requests.get(url, headers=headers, timeout=TIMEOUT, stream=True)
    if r.status_code == 304:
        LOG.info("%s export is unchanged (HTTP 304)", source or "source")
        return None, None
    r.raise_for_status()
    data = bytearray()
    for chunk in r.iter_content(1024 * 1024):
        if chunk:
            data.extend(chunk)
    return bytes(data), {
        "etag": r.headers.get("ETag"),
        "last_modified": r.headers.get("Last-Modified"),
    }


def _save_source_state(source, url, metadata):
    if not source or not DSN or not metadata:
        return
    etag = metadata.get("etag")
    last_modified = metadata.get("last_modified")
    if not etag and not last_modified:
        return
    with psycopg.connect(DSN) as conn:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS drug_source_state ("
            "source VARCHAR(32) PRIMARY KEY, url TEXT NOT NULL, etag TEXT, "
            "last_modified TEXT, updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW())"
        )
        conn.execute(
            "INSERT INTO drug_source_state(source,url,etag,last_modified) "
            "VALUES(%s,%s,%s,%s) "
            "ON CONFLICT(source) DO UPDATE SET "
            "url=EXCLUDED.url, etag=EXCLUDED.etag, "
            "last_modified=EXCLUDED.last_modified, updated_at=NOW()",
            (source, url, etag, last_modified),
        )
        conn.commit()


def _ensure_drug_columns(conn):
    """Миграция расширенных полей ГРЛС для уже существующих БД."""
    columns = {
        "registration_date": "DATE",
        "expiry_date": "DATE",
        "cancellation_date": "DATE",
        "production_stages": "TEXT",
        "pharmacotherapeutic_group": "VARCHAR(500)",
        "essential_drug": "BOOLEAN",
        "contains_controlled_substances": "BOOLEAN",
        "orphan_status": "VARCHAR(255)",
    }
    existing = {
        row[0] for row in conn.execute(
            "SELECT column_name FROM information_schema.columns "
            "WHERE table_schema=current_schema() AND table_name='drugs'"
        ).fetchall()
    }
    for name, sql_type in columns.items():
        if name not in existing:
            conn.execute(f"ALTER TABLE drugs ADD COLUMN {name} {sql_type}")


def _parse_date(value):
    value = _text(value, 32)
    if not value:
        return None
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Некорректная дата ГРЛС: {value}")


def _parse_bool(value):
    value = _text(value, 64)
    if not value:
        return None
    normalized = value.lower().replace("ё", "е")
    if normalized in {"да", "есть", "имеется", "1", "true", "yes", "присутствует"}:
        return True
    if normalized in {"нет", "отсутствует", "0", "false", "no"}:
        return False
    return None


def _start_run(conn, source):
    # Безопасная миграция для уже существующего PostgreSQL volume:
    # docker-entrypoint-initdb.d выполняется только при первом создании БД.
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS drug_gtin_pending (
            id BIGSERIAL PRIMARY KEY,
            gtin VARCHAR(14) NOT NULL,
            reg_number VARCHAR(50) NOT NULL,
            package_desc VARCHAR(500),
            first_seen_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            last_seen_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
            UNIQUE(gtin, reg_number)
        )
        """
    )
    conn.execute(
        "CREATE INDEX IF NOT EXISTS idx_drug_gtin_pending_reg ON drug_gtin_pending(reg_number)"
    )
    row = conn.execute(
        "INSERT INTO drug_import_runs(source,status) VALUES(%s,'running') RETURNING id",
        (source,),
    ).fetchone()
    conn.commit()
    return row[0]


def _finish_run(conn, run_id, status, seen, loaded, skipped, error=None):
    conn.execute(
        """
        UPDATE drug_import_runs
        SET finished_at=NOW(), status=%s, records_seen=%s,
            records_loaded=%s, records_skipped=%s, error_text=%s
        WHERE id=%s
        """,
        (status, seen, loaded, skipped, error, run_id),
    )
    conn.commit()


def _find_header_row(ws, required_aliases, max_rows=30):
    for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True), start=1):
        if _find_column(row, required_aliases) is not None:
            return row_no, row
    return None, None


def _grls_workbooks(data, root):
    archive = root / "grls.bin"
    archive.write_bytes(data)
    if zipfile.is_zipfile(archive):
        zroot = root / "unzipped"
        zroot.mkdir()
        with zipfile.ZipFile(archive) as z:
            candidates = [n for n in z.namelist()
                          if n.lower().endswith(".xlsx") and not n.endswith("/")]
            if not candidates:
                raise ValueError("В ZIP ГРЛС не найден XLSX")
            paths = []
            for name in candidates:
                target = zroot / Path(name).name
                with z.open(name) as src, target.open("wb") as dst:
                    dst.write(src.read())
                paths.append(target)
            return paths
    return [archive]


def _infer_grls_status(path):
    name = path.stem.lower().replace("_", " ")
    if "не действует" in name or "архив" in name:
        return "архив"
    if "действует" in name:
        return "действует"
    return None


def load_grls(data):
    """Загружает все XLSX из ZIP ГРЛС, не прерывая импорт из-за отдельных строк."""
    seen = loaded = skipped = 0
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        paths = _grls_workbooks(data, root)
        with psycopg.connect(DSN) as conn:
            _ensure_drug_columns(conn)
            conn.commit()
            run_id = _start_run(conn, "grls")
            try:
                for xlsx in paths:
                    wb = load_workbook(xlsx, read_only=True, data_only=True)
                    try:
                        for ws in wb.worksheets:
                            headers_row, headers = _find_header_row(
                                ws, HEADER_ALIASES["reg_number"], max_rows=40
                            )
                            if headers_row is None:
                                continue
                            rows = ws.iter_rows(min_row=headers_row + 1, values_only=True)
                            idx = {k: _find_column(headers, aliases)
                                   for k, aliases in HEADER_ALIASES.items()}
                            if idx["reg_number"] is None or idx["trade_name"] is None:
                                continue
                            inferred_status = _infer_grls_status(xlsx)
                            for row in rows:
                                seen += 1
                                try:
                                    reg = _text(row[idx["reg_number"]], 50)
                                    name = _text(row[idx["trade_name"]], 255)
                                    if not reg or not name:
                                        skipped += 1
                                        continue
                                    values = {
                                        k: _text(row[i], 500 if k in ("manufacturer", "holder", "production_stages") else 255)
                                        if i is not None and i < len(row) else None
                                        for k, i in idx.items() if k != "reg_number"
                                    }
                                    status = values.get("status") or inferred_status
                                    conn.execute(
                                        """
                                        INSERT INTO drugs(
                                            reg_number,trade_name,inn,dosage_form,dosage_value,
                                            manufacturer,holder,registration_date,expiry_date,
                                            cancellation_date,production_stages,pharmacotherapeutic_group,
                                            essential_drug,contains_controlled_substances,orphan_status,
                                            status,updated_at
                                        )
                                        VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,NOW())
                                        ON CONFLICT(reg_number) DO UPDATE SET
                                          trade_name=EXCLUDED.trade_name, inn=EXCLUDED.inn,
                                          dosage_form=EXCLUDED.dosage_form,
                                          dosage_value=EXCLUDED.dosage_value,
                                          manufacturer=EXCLUDED.manufacturer,
                                          holder=EXCLUDED.holder,
                                          registration_date=EXCLUDED.registration_date,
                                          expiry_date=EXCLUDED.expiry_date,
                                          cancellation_date=EXCLUDED.cancellation_date,
                                          production_stages=EXCLUDED.production_stages,
                                          pharmacotherapeutic_group=EXCLUDED.pharmacotherapeutic_group,
                                          essential_drug=EXCLUDED.essential_drug,
                                          contains_controlled_substances=EXCLUDED.contains_controlled_substances,
                                          orphan_status=EXCLUDED.orphan_status,
                                          status=EXCLUDED.status,
                                          updated_at=NOW()
                                        """,
                                        (
                                            reg, name, values.get("inn"), values.get("dosage_form"),
                                            values.get("dosage_value"), values.get("manufacturer"),
                                            values.get("holder"), _parse_date(values.get("registration_date")),
                                            _parse_date(values.get("expiry_date")),
                                            _parse_date(values.get("cancellation_date")),
                                            values.get("production_stages"),
                                            values.get("pharmacotherapeutic_group"),
                                            _parse_bool(values.get("essential_drug")),
                                            _parse_bool(values.get("contains_controlled_substances")),
                                            values.get("orphan_status"), status,
                                        ),
                                    )
                                    loaded += 1
                                    if loaded % BATCH_SIZE == 0:
                                        conn.commit()
                                except Exception as exc:
                                    skipped += 1
                                    LOG.warning("GRLS row %s skipped: %s", seen, exc)
                    finally:
                        wb.close()
                reconcile_pending(conn)
                conn.commit()
                _finish_run(conn, run_id, "success", seen, loaded, skipped)
            except Exception as exc:
                conn.rollback()
                _finish_run(conn, run_id, "error", seen, loaded, skipped, str(exc)[:4000])
                raise
    return seen, loaded, skipped

def _decode_mdlp(data):
    """Декодирует CSV без потери русских названий колонок."""
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8-sig", errors="replace")


def _csv_reader(data):
    """Возвращает reader и найденную строку заголовка МДЛП."""
    text = _decode_mdlp(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"

    rows = csv.reader(io.StringIO(text), dialect)
    headers = None
    for row in rows:
        if not any(str(value or "").strip() for value in row):
            continue
        if _find_column(row, GTIN_ALIASES) is not None and _find_column(row, REG_ALIASES) is not None:
            headers = row
            break

    if headers is None:
        raise ValueError("В CSV МДЛП не найдены заголовки GTIN и номера РУ")
    return headers, rows


def load_mdlp(data):
    """Загружает CSV МДЛП GTIN↔РУ.

    Поддерживает BOM, UTF-8/CP1251, разделители ;/,/TAB и служебные строки
    перед заголовком.
    """
    seen = loaded = skipped = 0
    with psycopg.connect(DSN) as conn:
        run_id = _start_run(conn, "mdlp")
        try:
            headers, reader = _csv_reader(data)
            gtin_i = _find_column(headers, GTIN_ALIASES)
            reg_i = _find_column(headers, REG_ALIASES)
            package_i = _find_column(headers, PACKAGE_ALIASES)
            if gtin_i is None or reg_i is None:
                raise ValueError("В CSV МДЛП не найдены GTIN и номер РУ")

            for row in reader:
                seen += 1
                try:
                    gtin = _normalise_gtin(row[gtin_i] if gtin_i < len(row) else "")
                    reg = _text(row[reg_i] if reg_i < len(row) else "", 50)
                    if not reg:
                        skipped += 1
                        continue
                    package = _text(row[package_i], 500) if package_i is not None and package_i < len(row) else None
                    drug = conn.execute("SELECT id FROM drugs WHERE reg_number=%s", (reg,)).fetchone()
                    if not drug:
                        # Публикации МДЛП и ГРЛС могут приходить в разном порядке.
                        # Не теряем корректную строку: складываем её во временную
                        # очередь и привяжем к drugs при следующем запуске.
                        conn.execute(
                            """
                            INSERT INTO drug_gtin_pending(gtin,reg_number,package_desc,last_seen_at)
                            VALUES(%s,%s,%s,NOW())
                            ON CONFLICT(gtin,reg_number) DO UPDATE SET
                              package_desc=EXCLUDED.package_desc,
                              last_seen_at=NOW()
                            """,
                            (gtin, reg, package),
                        )
                        skipped += 1
                        continue
                    conn.execute(
                        """
                        INSERT INTO drug_gtins(gtin,drug_id,package_desc)
                        VALUES(%s,%s,%s)
                        ON CONFLICT(gtin,drug_id) DO UPDATE SET package_desc=EXCLUDED.package_desc
                        """,
                        (gtin, drug[0], package),
                    )
                    loaded += 1
                    if loaded % BATCH_SIZE == 0:
                        conn.commit()
                except Exception as exc:
                    skipped += 1
                    LOG.warning("MDLP row %s skipped: %s", seen, exc)
            reconcile_pending(conn)
            conn.commit()
            _finish_run(conn, run_id, "success", seen, loaded, skipped)
        except Exception as exc:
            conn.rollback()
            _finish_run(conn, run_id, "error", seen, loaded, skipped, str(exc)[:4000])
            raise
    return seen, loaded, skipped


def reconcile_pending(conn):
    """Привязывает ранее неразрешённые GTIN к уже загруженным записям ГРЛС."""
    rows = conn.execute(
        "SELECT gtin, reg_number, package_desc FROM drug_gtin_pending ORDER BY id"
    ).fetchall()
    resolved = 0
    for gtin, reg_number, package_desc in rows:
        drug = conn.execute("SELECT id FROM drugs WHERE reg_number=%s", (reg_number,)).fetchone()
        if not drug:
            continue
        conn.execute(
            """
            INSERT INTO drug_gtins(gtin,drug_id,package_desc)
            VALUES(%s,%s,%s)
            ON CONFLICT(gtin,drug_id) DO UPDATE SET package_desc=EXCLUDED.package_desc
            """,
            (gtin, drug[0], package_desc),
        )
        conn.execute(
            "DELETE FROM drug_gtin_pending WHERE gtin=%s AND reg_number=%s",
            (gtin, reg_number),
        )
        resolved += 1
    if resolved:
        LOG.info("MDLP pending mappings resolved: %s", resolved)
    return resolved


def main():
    p = argparse.ArgumentParser()
    p.add_argument("source", choices=("grls", "mdlp"))
    p.add_argument("--file", help="локальный файл вместо скачивания")
    args = p.parse_args()
    if not DSN:
        raise SystemExit("DRUG_DB_DSN is required")
    url = GRLS_URL if args.source == "grls" else MDLP_URL
    if args.file:
        data = Path(args.file).read_bytes()
        metadata = None
    else:
        data, metadata = _download(url, args.source)
        if data is None:
            return

    result = load_grls(data) if args.source == "grls" else load_mdlp(data)
    _save_source_state(args.source, url, metadata)
    LOG.info("%s import complete: seen=%s loaded=%s skipped=%s", args.source, *result)


if __name__ == "__main__":
    main()
